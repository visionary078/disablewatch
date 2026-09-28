# -*- coding: utf-8 -*-
"""Generate the full glasses product plan as a Word document."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).with_name("see-next-step-glasses-plan-v1.docx")
NAVY = RGBColor(0x1E, 0x3A, 0x5F)
ACCENT = RGBColor(0xB4, 0x53, 0x09)
MUTED = RGBColor(0x4B, 0x55, 0x63)


def east_asia(run, name="微软雅黑"):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")


def set_style_font(style, size, bold=False, color=None, name="微软雅黑"):
    font = style.font
    font.size = Pt(size)
    font.bold = bold
    font.name = "Calibri"
    if color:
        font.color.rgb = color
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def p(doc, text, size=12, bold=False, color=None, align=None, space_after=8, space_before=0):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    pf = para.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = para.add_run(text)
    east_asia(run)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return para


def quote(doc, text):
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.left_indent = Cm(0.8)
    pf.space_after = Pt(12)
    run = para.add_run(text)
    east_asia(run)
    run.font.size = Pt(12)
    run.italic = True
    run.font.color.rgb = NAVY


def heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def table(doc, headers, rows, col_widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.autofit = True
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        para = hdr[i].paragraphs[0]
        run = para.add_run(h)
        east_asia(run)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(hdr[i], "1E3A5F")
    for r_i, row in enumerate(rows):
        cells = t.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            para = cells[c_i].paragraphs[0]
            run = para.add_run(str(val))
            east_asia(run)
            run.font.size = Pt(9.5)
            if r_i % 2 == 1:
                shade(cells[c_i], "F4F6F8")
    if col_widths:
        for row in t.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        para.clear()
        run = para.add_run(item)
        east_asia(run)
        run.font.size = Pt(12)


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.4)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run("看见下一步 · 眼镜产品策划案 v1.1  ·  内部讨论稿  ·  2026-09-24")
        east_asia(run)
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED

    styles = doc.styles
    set_style_font(styles["Normal"], 12, color=RGBColor(0x1F, 0x29, 0x37))
    styles["Normal"].paragraph_format.line_spacing = 1.35
    set_style_font(styles["Heading 1"], 18, True, NAVY)
    set_style_font(styles["Heading 2"], 14, True, NAVY)
    set_style_font(styles["Heading 3"], 12, True, ACCENT)

    # Cover
    p(doc, "看见下一步", 14, True, ACCENT, WD_ALIGN_PARAGRAPH.CENTER, 4, 48)
    p(doc, "眼镜产品策划案", 32, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, 6)
    p(
        doc,
        "3D 打印镜架  ·  挂脖电源  ·  摄像头 + 测距  ·  同一套识别网关",
        13,
        False,
        MUTED,
        WD_ALIGN_PARAGRAPH.CENTER,
        18,
    )
    quote(doc, "朝哪看，说哪边；近了请你停。手杖仍在你手里。")
    table(
        doc,
        ["项目", "内容"],
        [
            ["产品线", "小程序「爱总会看见」的第二形态：第一视角眼镜"],
            ["文档版本", "v1.1"],
            ["日期", "2026-09-24"],
            ["软件现状", "微信小程序 + FastAPI /infer + /glasses 网页"],
            ["线上网关", "https://watchapi.divesee.com"],
            ["当前主视觉", "小米 MiMo mimo-v2.5（现在这一档，不写死）"],
            ["换模型原则", "改网关 profile；眼镜、JSON 契约、清洗层不动"],
            ["长期记忆", "记忆张量 MemOS；短时主任务仍走现网关"],
            ["原则", "脸上只留传感器；不报厘米数；不替代导盲杖；不过马路"],
            ["配套文件夹", "仓库根目录 glasses/"],
        ],
    )

    heading(doc, "目录", 1)
    toc = [
        "1  一句话与结论",
        "2  背景：为什么现在做眼镜",
        "3  产品定义与边界",
        "4  对标：别人已经做过的",
        "5  推荐形态：挂脖眼镜",
        "6  重量、续航与充电",
        "7  传感器：摄像头、ToF、雷达",
        "8  软件如何接到现有网关",
        "9  视觉模型可替换：MiMo 只是现在",
        "10  长期记忆：与记忆张量 MemOS 合并",
        "11  客户交流话术",
        "12  三期路线、成本与验收",
        "13  风险",
        "14  近期行动",
    ]
    for line in toc:
        p(doc, line, 12, False, NAVY, space_after=4)

    heading(doc, "1  一句话与结论", 1)
    p(
        doc,
        "小程序解决「零硬件、扫码即用」；眼镜解决「举着手机对不准、腾不出手」。识别走云端网关，眼镜只负责拍、听、震、播。看图模型现在是 MiMo，后期换成更好的多模态，不改镜架。记得这个人，靠记忆张量，不靠把记忆模型戴到脸上。",
    )
    p(doc, "本策划的硬结论：", 12, True)
    bullets(
        doc,
        [
            "做挂脖电源环 + 可翻上镜架。相机必须在眼位，不要把主摄像头改挂胸口。",
            "脸上重量一期 ≤55 g，硬顶 80 g；2000–3000 mAh 电池放后颈。",
            "方位靠摄像头和云端多模态；测距靠 ToF / 60 GHz 雷达，对用户只说一臂内 / 较近 / 较远。",
            "MiMo 是当前主视觉，不是产品定义。换模型只改网关 profile，输出仍是六向方位 + 粗远近 + 一句口语。",
            "记忆张量 MemOS 做跨天偏好和纠正；现有 30 分钟主任务记忆保留。记忆不能盖过当前画面，也不能盖过风险。",
            "磁吸线一物两用：边充边用，掉了挂在身上，刮到门要自己断开。",
            "太阳能和走动发电补不上工作功耗，不写进卖点。",
            "对客户先说能帮什么，边界放后半句。",
        ],
    )

    heading(doc, "2  背景：为什么现在做眼镜", 1)
    p(
        doc,
        "试用里失败的不是「模型不会看图」，是采集姿态：货架前要举着手机，手被占住，镜头也不是眼睛朝向。社恐用户更不愿当街举起手机。",
    )
    p(doc, "眼镜不重新发明识别，只改三件事：", 12, True)
    bullets(
        doc,
        [
            "光轴跟着视线，方位才有意义。",
            "双手可以拿杖、拿购物篮。",
            "ToF / 雷达才能对准正前方，距离不再靠手机 VisionKit 猜。",
        ],
    )
    p(
        doc,
        "明确不做成 Envision 那种万元封闭整机，不把树莓派 5 戴在鼻梁上，不把研究原型的 400 g 头盔说成可穿戴产品。",
    )

    heading(doc, "3  产品定义与边界", 1)
    heading(doc, "3.1 给谁", 2)
    p(
        doc,
        "全盲、低视力、临时视觉受限；已经能用小程序，或愿意戴一副其貌不扬的打印镜架。社恐用户优先：不要闪光灯，不要胸前大支架。",
    )
    heading(doc, "3.2 一副眼镜只承诺", 2)
    bullets(
        doc,
        [
            "第一视角拍照，接现有 /infer：实时约 2.5 秒一帧、再看一眼、语音改主任务。",
            "正前方粗距离：一臂内 / 较近 / 较远 / 无法判断。",
            "高风险震动。口语例如：「正在帮你找可乐，正前方货架中部较近。」",
            "按住镜腿或按住说话键说出要找什么。",
        ],
    )
    heading(doc, "3.3 仍然不承诺", 2)
    bullets(
        doc,
        [
            "替代导盲杖、导盲犬、无障碍设施或人工协助。",
            "过马路决策。",
            "米级、厘米级测距播报。",
            "SLAM、连续 1080p 录像、全天不充电。",
        ],
    )

    heading(doc, "4  对标：别人已经做过的", 1)
    p(
        doc,
        "要专用产品的任务、消费级的重量和充电、开源的拆分架构。不抄万元封闭机，也不抄 400 g 论文头盔。",
    )
    heading(doc, "4.1 专用视障 / 低视力", 2)
    table(
        doc,
        ["产品", "形态", "重量", "电池 / 充电", "价格", "和我们的关系"],
        [
            [
                "Envision Glasses",
                "Google Glass EE2 + 钛架，读文字 / 找物",
                "约 46 g",
                "820 mAh，USB-C，5–6 h",
                "Home 约 $2499",
                "重量上限的标杆。贵在软件和认证。",
            ],
            [
                "OrCam MyEye 3 Pro",
                "磁吸夹在自有镜框，13 MP",
                "模组 22.5 g",
                "320 mAh，磁吸约 1 h 充满，续航 60–90 min",
                "约 $4250",
                "夹具比整机轻；鼻梁塞大电池会牺牲续航。",
            ],
            [
                "eSight Go",
                "低视力放大显示",
                "约 170 g",
                "约 6 h",
                "约 $4950",
                "反面：过 100 g 全天戴会劝退。我们不做镜片放大。",
            ],
            [
                "NuEyes Pro 4",
                "Android 眼镜 + OCR",
                "约 102 g，电源常外挂",
                "线供电",
                "$3895 起",
                "外挂电池可借鉴，整机太重不学。",
            ],
        ],
    )
    heading(doc, "4.2 消费级 AI 眼镜（抄重量和充电，不抄任务）", 2)
    p(
        doc,
        "普通眼镜 20–30 g，国内横评里 AI 眼镜 36–53 g。能卖出去的产品，镜身只有 150–250 mAh，一天靠充电盒。",
    )
    table(
        doc,
        ["产品", "重量", "镜身电池", "充电", "可抄"],
        [
            ["小米 AI 眼镜", "约 40 g", "200 mAh 级", "Type-C 直插", "充电简单；盲操找口不一定好。"],
            [
                "雷鸟 V3",
                "约 39 g",
                "159 mAh",
                "磁吸 + 3000 mAh 盒 + 约 115 cm 边充边用线",
                "最值得抄：小电池、盒子、口袋里的线。",
            ],
            ["Rokid Glasses", "约 49–52 g", "约 210 mAh", "磁吸约 30 min 充满，有盒", "磁吸比 Type-C 孔更好摸。"],
            ["Ray-Ban Meta", "约 53 g", "200 mAh 级", "充电盒", "社交接受度高；没有测距。"],
            ["Brilliant Labs Frame", "约 39 g", "210 mAh", "USB-C 底座另带 140 mAh", "开源、可 3D 打印周边。"],
            ["SGG LOOP / Rokid 充电环", "28–93 g 颈环", "450–900 mAh 挂脖", "磁吸或 USB-C", "证明颈挂供电已经有人卖。"],
        ],
    )
    heading(doc, "4.3 只做避障、开源与论文", 2)
    table(
        doc,
        ["项目", "做法", "启示"],
        [
            ["Sunu Band", "腕上超声波 40 g，约 14 h", "避障用触觉；眼镜上的 ToF/雷达也应能只震不说话。"],
            ["BuzzClip", "衣领超声波夹 62 g", "超重就把传感器改夹帽檐，不要加厚鼻梁。"],
            ["WeWALK 2", "智能杖手柄，地面仍交给杖", "眼镜不要抢杖的活。"],
            ["OpenGlass", "ESP32 夹在成镜上，< $25，手机算", "传感/计算拆开，和我们 /glasses 同构。"],
            ["PathSense", "帽檐 + Pi Zero + 多颗 ToF", "ToF 网格有效，电池在身上。"],
            ["宁波大学 2018", "3D 支架 + RealSense + TI 雷达", "雷达 + 视觉在光照变化时更稳。"],
            ["Nature Comm. 2025", "眼镜约 400 g（含 80 g 电池）", "论文能跑通 ≠ 能戴出门。"],
        ],
    )
    p(
        doc,
        "我们卡在中间：OpenGlass 的拆分 + Envision 的任务 + 雷鸟的充电习惯 + 挂脖环把电池和防掉放身上 + VL53/雷达补避障。",
        12,
        True,
    )

    heading(doc, "5  推荐形态：挂脖眼镜", 1)
    p(
        doc,
        "挂脖值得做，而且比裤袋长线更适合视障：眼镜始终在身上，摸得到、掉不远、电也在身上。只做对的那一种：相机仍在眼位，电池和绳子在脖子上。",
    )
    table(
        doc,
        ["形态", "做不做", "原因"],
        [
            [
                "A. 挂脖电源环 + 可翻上镜架",
                "做（主方案）",
                "鼻梁 50 g 级传感器；2000–3000 mAh 在后颈；摘下挂胸；磁吸既充电又防掉。",
            ],
            ["B. 眼镜链 / 纯绳子", "第 0 期可做", "防掉有效，不供电。一天就能验证挂胸好不好摸。"],
            [
                "C. 胸前摄像头当主眼",
                "不做主方案",
                "光轴在胸口，「左前方」不是脸朝的方向。胸前最多后来加过头障碍雷达。",
            ],
        ],
    )
    heading(doc, "5.1 三种使用状态", 2)
    p(doc, "状态 1 · 戴着看：镜架在脸上，两腿吸在挂脖两端。朝哪看，相机朝哪。挂脖供电，2.5 秒拍一帧。")
    p(
        doc,
        "状态 2 · 挂在胸前：一只手摘下镜架，磁吸还连着，眼镜垂在胸口。判定不在脸上后立刻停拍照、停播报。摸胸口就能戴回去，不用在包里找。",
    )
    p(doc, "状态 3 · 过夜：整圈套在 USB-C 座上，镜架仍吸在环上，零件不分开。")
    p(doc, "没有电源键。戴上并吸好 = 开始看；摘下挂胸 = 停下；长按左腿大键 = 说话。")
    heading(doc, "5.2 和裤袋磁吸线比", 2)
    table(
        doc,
        ["", "裤袋盒 + 1 m 线", "挂脖环"],
        [
            ["眼镜掉了", "能兜，但线在手杖半径里扫", "掉在胸口，路径短"],
            ["找得到吗", "盒在口袋，镜可能在桌上", "永远套在脖子上"],
            ["转头", "长线会扯衣角", "软管跟头动，12–18 cm"],
            ["社交", "一根线垂到裤袋，像在充电", "像挂脖耳机 / 工牌"],
            ["夏天", "较好", "后颈会热，要硅胶透气"],
            ["大衣 / 安全带", "较少干涉", "要能从衣领里掏出来"],
        ],
    )
    p(
        doc,
        "日常室内、超市、机构参观用挂脖。冬天厚外套、安全带场景保留「盒 + 短线」为备选，同一套磁吸口兼容。",
    )
    heading(doc, "5.3 机械尺寸（第一版按这个打）", 2)
    bullets(
        doc,
        [
            "后颈舱 78 × 42 × 20 mm，扁圆、贴后颈。PETG + 2 mm 硅胶垫。USB-C 朝下，喇叭口盲插。外侧一条竖棱表示正放，凸点表示电量。",
            "颈环：中空硅胶管外径 8–9 mm，总长可调 380–450 mm（颈围 36–42 cm）。两级卡扣，不要细扣。",
            "磁吸头：φ10–12 mm，长 18 mm，2–4 弹簧针。断开力 2.5–4 N，能吊住 55 g，刮到门自己掉。拉力走外壳，不进焊盘。",
            "线长挂脖侧只需 12–18 cm；若退回裤袋方案则用 70–90 cm，太长眼镜落地了线才绷紧。",
            "领夹夹在衣领，减少和导盲杖、购物篮缠绕。",
            "挂胸时镜头朝向身体，不要朝外扫路人。耳挂当挂钩扣在颈环上。",
        ],
    )
    heading(doc, "5.4 软件生死线：在不在脸上", 2)
    p(doc, "挂胸时相机对着衣服或地面，若继续识别会乱报方位。")
    table(
        doc,
        ["信号", "戴着", "挂胸"],
        [
            ["俯仰", "近水平 ±25°", "经常 >50° 朝下或朝里"],
            ["磁吸", "两端都通", "仍通（没断）"],
            ["接近（可选）", "鼻托 / 镜腿贴皮肤", "离开皮肤"],
        ],
    )
    bullets(
        doc,
        [
            "挂胸连续 1.5 秒 → 停直播、停播报、停雷达。",
            "戴回 0.8 秒 → 「开始看。」再出第一帧。",
            "磁吸断开且自由落体 → 「眼镜掉了。」若被线兜住，改口「还在身上。」",
            "session 和「找可乐」主任务挂胸期间保留，戴回去继续找。",
        ],
    )

    heading(doc, "6  重量、续航与充电", 1)
    heading(doc, "6.1 重量预算", 2)
    table(
        doc,
        ["部位", "第 1 期", "第 2 期", "硬上限"],
        [
            ["脸上（架、相机、线、小电池）", "≤ 55 g", "≤ 65 g（加 ToF）", "80 g，超过改帽檐"],
            ["后颈挂脖舱", "90–130 g", "同左", "先用 120 g 空壳戴两小时"],
            ["戴着时鼻梁", "42–55 g", "同左", "和 Envision 46 g 一个量级"],
        ],
    )
    table(
        doc,
        ["脸上零件", "约重", "备注"],
        [
            ["PETG 镜架 + 鼻托", "28–35 g", "镜腿加厚走线"],
            ["场景相机", "3–5 g", "B0066 针孔或 Module 3（4 g）"],
            ["ESP32-S3 或 UVC", "4–6 g", "只采图，不跑模型"],
            ["保命电芯 150–200 mAh", "4–6 g", "断线还能说「线掉了」"],
            ["磁吸触点、螺丝", "3–5 g", ""],
            ["合计", "约 42–57 g", "落在专用机重量附近"],
        ],
    )
    heading(doc, "6.2 功耗与续航", 2)
    p(
        doc,
        "连续 1080p 录像会在一小时内掏空 150–250 mAh。必须沿用小程序策略：每 2.5 秒抓一帧 JPEG，中间休眠。",
    )
    bullets(
        doc,
        [
            "工作平均大约 80–150 mA（约 0.3–0.5 W），不是听音乐那种待机。",
            "镜身 250 mAh 单独扛：货架巡视约 1.5–3 小时。",
            "后颈 2500–3000 mAh（约 7–11 Wh）：按 6–10 小时货架巡视设计，够一个工作日。不必上 7000 mAh 项圈。",
            "雷达必须占空比（例如 5 Hz 扫一次），不能一直 FMCW。",
        ],
    )
    heading(doc, "6.3 充电与防掉", 2)
    p(doc, "磁吸是第一优先：盲操靠「吸住了」，OrCam、雷鸟、Rokid 都是这条路。")
    bullets(
        doc,
        [
            "挂脖主路径：后颈舱 USB-C 过夜 2–3 小时充满；磁吸输出限流 0.5 A。",
            "备选裤袋盒：3000 mAh，30–45 分钟给镜身充满，70–90 cm 线边用边充。",
            "吸力 2.5–4 N：托住眼镜，门把手刮到必须自己断开，避免勒脖子。",
            "耳挂是防掉第一道，磁吸线是第二道兼充电。",
            "不要 Micro-USB。Type-C 直插镜腿只作工程样机。",
            "电量用语音和可摸凸点，不靠屏幕：开机报大约还能用多久；低于 20% 短震三次。",
        ],
    )
    heading(doc, "6.4 太阳能和走动发电", 2)
    p(
        doc,
        "补得上待机，补不上「正在看」。KIT 太阳能眼镜室内每只镜片约 200 微瓦，只够温度计。镜腿柔性电池户外乐观几十毫瓦，相对 0.4 W 工作功耗大约一成，超市里几乎为零。走动、体温是微瓦级。真正能加续航的是：2.5 秒一帧、充电盒 / 挂脖大电池、雷达占空比。太阳能若做，只考虑贴在充电盒盖上，不贴鼻梁。",
    )

    heading(doc, "7  传感器：摄像头、ToF、雷达", 1)
    p(
        doc,
        "可以做到「基本」的方位和测距，但不是同一颗传感器各干一半。摄像头认是什么、在哪一侧；雷达 / ToF 认正前方有多近、是不是在靠近。合起来是「可乐在左前方，较近」，不是「37 厘米、偏 18 度」。",
    )
    heading(doc, "7.1 推荐栈", 2)
    table(
        doc,
        ["位置", "器件", "朝向", "外壳"],
        [
            ["鼻梁略偏右", "RGB 主摄 IMX708 或 OV5647 针孔", "视线正前，俯角 5–8°", "光学开窗，PETG 不要挡镜头"],
            ["主摄正下方", "VL53L8CX（第 2 期）", "与相机共光轴", "透 940 nm 的深色窗"],
            ["右镜腿前段", "XM125 / A121（第 3 期）", "正前略外", "1–2 mm 塑料雷达罩，前方无金属"],
            ["左镜腿", "磁吸 + 按住说话键", "—", "键程要深"],
            ["后颈", "大电池", "—", "硅胶贴肤"],
        ],
    )
    heading(doc, "7.2 适合上镜的摄像头", 2)
    table(
        doc,
        ["型号", "接口", "尺寸 / 重量", "水平视场", "判定", "用途"],
        [
            ["Arducam B0066 OV5647 针孔", "MIPI", "颈部 6 mm · ~3 g", "62°", "首选上镜", "最瘦，藏进鼻梁"],
            ["Arducam B006603 M6", "MIPI", "~3 g", "72°", "首选上镜", "货架更不容易切边"],
            ["树莓派 Camera Module 3 IMX708", "MIPI", "25×24×11.5 mm · 4 g", "66°", "首选原型", "近距对焦看价格牌"],
            ["Module 3 Wide", "MIPI", "4 g", "102°", "广角避障", "不要当找可乐主摄"],
            ["Arducam B0304 IMX708 USB", "UVC", "转接板 38×38", "66°", "接现有眼镜页", "免驱进 /glasses"],
            ["USB 7–8 mm 1080p 针孔", "UVC", "镜头约 8 mm", "60–90°", "好装", "要抽检清晰度"],
            ["IMX219 M12 160–170°", "MIPI", "M12 座更大", "160°+", "不要当主摄", "鱼眼会把方位挤糊"],
            ["OV5640 / ESP32", "DVP", "常见 24×25", "约 66°", "最低成本", "只适合证明链路"],
            ["HQ / OAK-1 / D435", "USB / MIPI", "过重", "—", "不要上脸", "胸前或背包"],
        ],
    )
    heading(doc, "7.3 测距与雷达", 2)
    table(
        doc,
        ["型号", "原理", "距离 / 视场", "室外", "判定"],
        [
            ["ST VL53L8CX", "8×8 ToF", "室内 ~4 m，对角 65°", "明显缩短", "第一颗测距，货架和一臂内"],
            ["ST VL53L7CX", "8×8 ToF 更宽", "室内 ~3.5 m，90°", "差", "更宽避障"],
            ["Acconeer XM125 / A121", "60 GHz 相干雷达", "人约 7 m，约 65°", "几乎不受影响", "眼镜雷达首选：暗、玻璃门"],
            ["Infineon BGT60TR13C", "60 GHz 1TX3RX", "0.2–15 m，约 90°", "好", "要左右角再用，更费电"],
            ["HLK-LD2450", "24 GHz 跟人", "运动体 ~6 m", "室内为主", "货架当背景，不当主避障"],
            ["LD2410 / 超声波 / 旋转雷达", "存在或声呐或 360°", "各异", "各异", "不要"],
        ],
    )
    p(
        doc,
        "光轴要同心，否则雷达打到的「最近货架」不是相机说的那罐可乐。对用户仍然禁止播米、厘米、步。",
    )

    heading(doc, "8  软件如何接到现有网关", 1)
    bullets(
        doc,
        [
            "小程序与眼镜网页共用 /infer、/tts、/asr，密钥不出端。",
            "眼镜页：software/server/static/glasses/ ，地址 /glasses/?token=业务口令。",
            "实时约 2.5 秒一帧；上一帧没完就丢帧，避免过时画面。",
            "语音改主任务后，后续帧继续找同一目标，不退回避障套话。",
            "ToF / 雷达只改 distance_band 和风险，不改主任务。",
            "眼镜与小程序用不同 session_id，互不抢锁；配对后共用 user_id，长期记忆才能跟人走。",
            "镜上不跑大模型。换多模态、换记忆后端，都在网关完成。",
            "开机语音：朝前看，我说哪边、近不近。手杖照拿，过马路别靠我。",
        ],
    )

    heading(doc, "9  视觉模型可替换：MiMo 只是现在", 1)
    p(
        doc,
        "产品契约不是「必须用小米 MiMo」，而是「一张图进网关，出来一句能执行的下一步」。MiMo mimo-v2.5 是当前主路径上已经接好的一档。后期换 Qwen-VL、Gemini、豆包视觉或其它更强多模态，改的是 model_profiles.json，不是镜架、不是 /glasses 页、不是用户听到的句式。",
    )
    heading(doc, "9.1 为什么必须可换", 2)
    bullets(
        doc,
        [
            "看货架、读价签、暗光、玻璃门，各家多模态每年都在变；把厂商写进硬件等于把产品锁死。",
            "眼镜 2.5 秒一帧，时延、价格、中文口语随时可能让我们换家。",
            "网关已经按 profile 选模型：primary 现在是 MiMo，backup 是空位，fusion.py 里还有 verify 异构复核（桌面主路径尚未接线）。换模型是这条路的本意，不是事后补丁。",
        ],
    )
    heading(doc, "9.2 换模型时什么不许动", 2)
    table(
        doc,
        ["层", "稳定契约", "换模型时"],
        [
            ["眼镜 / 小程序", "拍 JPEG、带 session、带语音、播 speech", "不改固件、不改按键"],
            ["POST /infer", "image + question + mode + distance_band", "可多一个 model_profile"],
            ["模型输出", "intent / direction 六向 / proximity / risk / speech", "新模型必须被洗成这套 JSON"],
            ["清洗与融合", "禁止米厘米步；风险只升不降；不确定就少说", "任何新模型都要过这一关"],
            ["播报", "正在帮你找×，某方位，一臂内/较近/较远", "禁止变成看图作文"],
        ],
    )
    p(
        doc,
        "新模型若爱描写「货架上有许多瓶装饮料，灯光明亮」，网关必须挡掉。对人只留下一步。换得再好，也不许过马路、不许替代手杖。",
    )
    heading(doc, "9.3 下一档怎么选", 2)
    table(
        doc,
        ["门槛", "要求"],
        [
            ["时延", "precise 可到数秒；live 应能配合 2.5 秒丢帧，不能排队"],
            ["格式", "能稳定吐 JSON，或能被现有解析器收成 JSON"],
            ["中文口语", "短句，不像说明书，不像英文翻译腔"],
            ["货架", "罐装、价签、左右排，比风景描述重要"],
            ["听话", "有[主任务记忆]时必须找当前目标，不许改口报通道"],
            ["成本", "连续巡视可承受；比「榜单第一」优先"],
            ["密钥", "只在网关环境变量，不出小程序、不出眼镜页"],
        ],
    )
    p(
        doc,
        "候选只作后期选项，不写进对外承诺：通义 Qwen-VL、豆包视觉、Gemini、GPT 视觉、开源 VL 自建。谁上 backup / 谁上 verify，用同一套货架样张打分，不听发布会。记忆张量的 Metis / 忆立方不当主视觉——那是记忆侧，看图仍用多模态视觉模型。",
    )
    heading(doc, "9.4 接线方式", 2)
    bullets(
        doc,
        [
            "改 model_profiles.json 的 primary，或先填 backup 做热切换。",
            "需要双跑时再打开 verify profile：主模型出方位，验证模型只抬风险、不抢播报。现在融合函数在，配置还没有接线，对外不可说已上线。",
            "眼镜请求可以不带模型名，由网关默认；评测时才指定 profile。",
            "换模型验收：同一组「找可乐 / 纠正左边 / 挂胸停拍」录音，speech 仍短，方向不乱跳。",
        ],
    )

    heading(doc, "10  长期记忆：与记忆张量 MemOS 合并", 1)
    p(
        doc,
        "现有 task_memory 只记得这一趟、这 30 分钟要找什么。眼镜整天戴着，没有跨天记忆就会每天从零自我介绍。记忆张量（MemTensor）的 MemOS 补的是「这个人」：常买什么、怎么叫那只杯子、上次哪里说错了。合的是记忆层，不是把他们的记忆模型换成眼镜的眼睛。看图模型换成哪一家，记忆层都还在。",
    )
    heading(doc, "10.1 两层记忆不要揉成一层", 2)
    table(
        doc,
        ["层", "现在 / 计划", "寿命", "作用"],
        [
            ["短时主任务", "已有 services/task_memory.py", "30 分钟，最近 8 句", "说完找可乐，下一帧还在找"],
            ["长期用户记忆", "MemOS Cloud 或自建开源 MemOS", "跨天、跨设备", "默认无糖；药在第二层；忘掉就能忘"],
        ],
    )
    heading(doc, "10.2 合什么、不合什么", 2)
    p(doc, "合：偏好、别名、播报要短、纠正、经用户同意的弱场景线索（这家店上次可乐在右前）。", 12, True)
    p(
        doc,
        "不合：不替换 /infer 视觉模型；不把原图、人脸、连续视频写入记忆库；不让记忆盖过风险；不在播报里念「根据您的历史偏好」；不把 MemOS 部署到镜架上。",
    )
    heading(doc, "10.3 架构", 2)
    p(
        doc,
        "小程序 openid 与眼镜扫码配对成同一个 user_id。无 user_id 时行为与现在完全一致。检索只在新语音、主任务 new/switch、或本趟第一次拉画像时发生；超时 250–400 ms 则丢掉长期记忆、照常看图。写入在返回播报之后异步做。",
    )
    p(
        doc,
        "Prompt 只多一小段 [用户长期记忆]，并写死：禁止复述本段；禁止用记忆代替当前画面。记忆说在右、镜头里没有，就必须改口。",
    )
    heading(doc, "10.4 眼镜上长什么样", 2)
    table(
        doc,
        ["场景", "没有长期记忆", "合并之后"],
        [
            ["进超市", "每次说帮我找可乐", "「找可乐」默认无糖罐装"],
            ["在家找药", "每次解释柜子", "常用药第二层作先验，仍以画面为准"],
            ["纠正过一次", "下次还错", "同一类目标先避开旧错误"],
            ["戴眼镜出门", "和小程序是两个人", "扫码配对后是同一个人"],
        ],
    )
    quote(doc, "正在帮你找无糖可乐，右前方货架中部，比较近。")
    heading(doc, "10.5 隐私", 2)
    bullets(
        doc,
        [
            "只存结构化短句，默认不存图。",
            "眼镜未配对 = 匿名会话，不写长期库。",
            "用户能说忘掉可乐、忘掉今天。家属账号与本人分开。",
            "网关自包一层 LongMemory 接口，MemOS 只是其中一个后端，避免厂商锁定。",
            "Cloud 试用前先写清数据驻留和删除凭证；密钥只放网关。",
        ],
    )
    heading(doc, "10.6 记忆侧三期", 2)
    table(
        doc,
        ["期", "做什么", "验收"],
        [
            [
                "0",
                "user_id + 配对；偏好/别名；检索只在语音与任务切换；超时跳过",
                "无密钥时与现网一致；有记忆时第一次「找可乐」即带无糖；实时帧不明显变慢",
            ],
            ["1", "纠正写入；可选场景事件；忘掉××", "同一错误连续两天不重复；能清掉一条"],
            ["2", "小程序与眼镜打通；可自建 MemOS", "换设备不必重新教偏好"],
        ],
    )
    p(doc, "不做：每人一份 LoRA、用 Metis 当主视觉、每帧全库检索。")

    heading(doc, "11  客户交流话术", 1)
    p(doc, "原则：先说他能听见什么、手能做什么；边界放后半句。不讲二维空间、网关、FMCW、SLAM、APP_TOKEN。", 12, True)
    heading(doc, "11.1 三十秒", 2)
    quote(
        doc,
        "戴着眼镜或举着手机，朝哪看，我就说哪边、远不远、要不要先停。你说「帮我找可乐」，后面一直找可乐。手杖还是要拿着。我是旁边帮看的人，不是替你走路的人。手机扫码就能用。眼镜是为了腾出手。",
    )
    heading(doc, "11.2 挂脖怎么说", 2)
    quote(
        doc,
        "眼镜戴在眼睛上，电池挂在脖子上。摘下来就挂在胸前，摸得到，摔不着。吸上就能边走边充。朝哪看，说哪边。近了请你停。手杖还在你手里。",
    )
    heading(doc, "11.3 问答对照", 2)
    table(
        doc,
        ["客户原话", "别这样接", "这样接"],
        [
            ["能测距吗？", "毫米波精度可达毫米。", "能告诉你近不近。一臂之内请你先停。不报 37 厘米。"],
            ["能知道方位吗？", "二维空间六向枚举。", "左、左前、正前、右前、右。听完手知道往哪侧偏。"],
            ["有雷达就能出门？", "可以避障导航。", "正前方有东西靠近我会提醒。地面、马路还靠手杖。"],
            ["和豆包有什么区别？", "多模态行动辅助系统。", "豆包描写画面。我们说下一步。"],
            [
                "和 Envision / 小米眼镜呢？",
                "他们没有我们的融合算法。",
                "Envision 完整但贵。小米轻，但不说货架下一步。我们先把听得懂的下一步做好。",
            ],
            [
                "你们用哪个模型？",
                "我们用小米 MiMo，以后再说。",
                "现在用一档看图模型。以后换成更好的，你听到的还是那句下一步。模型名不是卖点。",
            ],
            [
                "会不会越用越懂我？",
                "我们接入了记忆操作系统。",
                "你常买什么、怎么叫家里那只杯子，说过一次我可以记住。你说忘掉就忘掉。",
            ],
            ["准吗？", "准确率 95%。", "看不清就说还在找。宁可少说，不带错路。"],
            ["能过马路吗？", "高风险会提示。", "不能。过马路不要用这个。"],
            ["多重？能戴多久？", "我们做了轻量化设计。", "脸上五十克上下。电池在脖子上。掉了挂着，不摔地上。"],
        ],
    )
    heading(doc, "11.4 产品里听得到的句子", 2)
    table(
        doc,
        ["场景", "不要像说明书", "对人说"],
        [
            ["欢迎", "本软件只提供环境信息辅助…", "朝前看，我说哪边、近不近。手杖照拿，过马路别靠我。"],
            ["开始", "戴上，开始辅助", "开始看"],
            ["语音", "按住说话说出需求", "按住，说要找什么"],
            ["确认", "精确确认", "再看一眼"],
            ["口令", "业务口令", "使用密码"],
            ["没看见", "未能确认目标，请重新拍摄", "还在帮你找。眼前是货架。"],
            ["电量", "电量 18%", "电不多了，该给后颈充电了"],
        ],
    )
    heading(doc, "11.5 演示五步", 2)
    bullets(
        doc,
        [
            "「我先让它自己看。」",
            "「现在我说：帮我找矿泉水。」",
            "「你听：它还在找同一瓶，没有改口说通道。」",
            "靠近一点：应说更近或震一下。",
            "收尾：「手杖我没放下。它只是多了一双眼睛。」失败了就说「这句没看清，我们设计成不编。」不要解释状态码。",
        ],
    )

    heading(doc, "12  三期路线、成本与验收", 1)
    table(
        doc,
        ["期", "脸上", "挂脖 / 电", "验收"],
        [
            [
                "0（3 天）",
                "任意成镜",
                "买硅胶眼镜链，验证挂胸",
                "闭眼能摸到胸前的眼镜",
            ],
            [
                "1（4–6 周）",
                "针孔或 Module 3，接 /glasses",
                "后颈 2500 mAh + 双磁吸供电",
                "戴着走货架，双手空，能找可乐。脸上 ≤55 g。掉了不碰地。",
            ],
            [
                "2（+3–4 周）",
                "VL53L8CX 共光轴",
                "领夹、IMU 挂胸检测、凸点电量",
                "一臂内能先停；挂胸绝对不再播方位。",
            ],
            [
                "3（+4 周）",
                "镜腿 60 GHz；超 80 g 则雷达改后舱",
                "同一环不加重",
                "暗通道、玻璃门不瞎。重量仍 ≤80 g。",
            ],
        ],
    )
    heading(doc, "12.1 打样成本（人民币大约）", 2)
    table(
        doc,
        ["项", "约"],
        [
            ["打印件 + 硅胶管 + 磁吸头", "80–150"],
            ["2500 mAh 电芯 + 充放电板", "40–80"],
            ["镜架相机 / ESP32", "80–200"],
            ["五金、线", "20"],
            ["第 1 期合计", "约 220–450"],
            ["第 2 期加 ToF", "+40–80"],
            ["第 3 期加 XM125", "+150–300"],
        ],
    )
    p(
        doc,
        "对比：Envision 两万元人民币量级，OrCam 三万以上，雷鸟 V3 约 1800 但不说货架下一步。打样成本不是售价。",
    )
    heading(doc, "12.2 打印件清单（第 1 期）", 2)
    bullets(
        doc,
        [
            "nape_shell_a / b.stl 后颈舱对开",
            "pogo_male_left / right.stl 软管端磁吸头",
            "pogo_female_temple.stl 镜腿插座",
            "frame_v1.stl 镜架，鼻梁相机开窗 φ8",
            "collar_clip.stl、earhook.stl、usb_flare.stl",
            "材料：PETG；贴肤用买来的硅胶，不要硬塑料贴脖子。",
        ],
    )
    heading(doc, "12.3 戴 20 分钟验收", 2)
    bullets(
        doc,
        [
            "闭眼独立完成：戴上、吸上、说话、摘下挂胸、再戴上。",
            "故意松手，眼镜不得碰地。",
            "软管绕椅背，磁吸应先于脖子受力断开。",
            "挂胸时绝对不能再播「左前方有可乐」。",
            "后颈空壳 120 g 戴两小时，勒就退回裤袋盒。",
        ],
    )

    heading(doc, "12.4 软件三期（模型与记忆）", 2)
    table(
        doc,
        ["期", "模型", "记忆", "验收"],
        [
            [
                "与硬件第 1 期并行",
                "主路径仍 MiMo；补齐 backup 空位的配置规范",
                "user_id 可选；无密钥则零行为",
                "不带记忆、不换模型时，与现网完全一致",
            ],
            [
                "与硬件第 2 期并行",
                "用同一组货架样张评一档更好的多模态，通过则切 backup 或切 primary",
                "偏好/别名生效；纠正可写",
                "换模型后句式仍短；「找可乐」能带上无糖",
            ],
            [
                "更后",
                "需要时再接线 verify 异构复核",
                "小程序与眼镜记忆打通",
                "换设备不必重新教；双模型不得抢播报",
            ],
        ],
    )

    heading(doc, "13  风险", 1)
    table(
        doc,
        ["风险", "为什么会发生", "对策"],
        [
            ["超重压鼻", "线材电芯螺丝偷偷加到 90 g", "超 80 g 改帽檐；每周称一次"],
            ["后颈坠 / 热", "3000 mAh 贴皮肤", "先 120 g 空壳试戴；硅胶垫；温升 <8°C"],
            ["连续拍照把电打光", "开发时当摄像机", "默认 2.5 秒一帧；禁止 1080p 常录"],
            ["挂胸仍在识别", "IMU 没做", "第 2 期必须先做停拍，再加 ToF"],
            ["磁吸勒颈", "吸力做成死扣", "2.5–4 N，刮到必须断开"],
            ["雷达当导盲", "演示话说满", "开机那句不变；与杖共用"],
            ["太阳能写成卖点", "客户问续航", "只说盒子/挂脖充电，不承诺晒太阳就能用"],
            ["PLA 热变形", "夏天软", "PETG；电芯隔离"],
            ["换模型变成长描写", "新多模态爱看图作文", "清洗层挡掉；验收仍是一句下一步"],
            ["把产品写成某厂模型", "对外点名 MiMo", "对外说看图模型可换；对内才写当前 profile"],
            ["记忆幻觉", "MemOS 说在右，画面没有", "画面优先；记忆只做 hint"],
            ["记忆拖死实时帧", "每帧全库检索", "只在语音/切换时检索；超时跳过"],
            ["眼镜借给别人", "仍用上一人的药柜记忆", "摘下超时；长按解除配对"],
        ],
    )

    heading(doc, "14  近期行动", 1)
    bullets(
        doc,
        [
            "量使用者颈围、耳到锁骨、鼻梁到后颈。",
            "用电子秤称空镜架、相机、电芯，填进重量表，不要靠估。",
            "打后颈空壳 + 硅胶管，装沙袋到 120 g，戴两小时。",
            "磁吸头：吊 55 g 不脱，吊 400 g 应脱开。",
            "通 5 V 只供电，眼镜继续打开现有 /glasses 页。",
            "见客户前过一遍第 11 章话术。",
            "加 IMU 挂胸停拍之后，再谈 ToF 和雷达。",
            "网关 /infer 增加可选 user_id；无此字段行为不变。",
            "LongMemory 先做空操作后端，再接 MemOS；手工写入 5 条偏好用可乐用例试 Prompt。",
            "用同一组货架样张列一档备选多模态的打分表，不先改 primary。",
        ],
    )
    p(doc, "配套文件（仓库 glasses/ 文件夹）", 12, True, space_before=12)
    bullets(
        doc,
        [
            "hardware-plan-v1.md 硬件策划原文",
            "neckband-plan-v1.md 挂脖尺寸与结构",
            "customer-talking-points-v1.md 话术原文",
            "memos-long-memory-plan-v1.md 与记忆张量合并方案",
            "网页代码 software/server/static/glasses/",
        ],
    )
    p(
        doc,
        "本文是内部讨论稿 v1.1，不替代技术文档 v1.6 对已上线小程序的描述。当前线上主视觉仍是 MiMo mimo-v2.5；双模型验证视觉在桌面主路径仍未接线；MemOS 尚未接线。对外不可把后两件事说成已上线。",
        11,
        False,
        MUTED,
        space_before=18,
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
